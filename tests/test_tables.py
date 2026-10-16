import json

import pytest
from openpyxl import Workbook, load_workbook

from web_task_robot.config import JobError
from web_task_robot.tables import read_rows, write_rows


def test_read_csv_with_bom_and_blank_lines(tmp_path):
    p = tmp_path / "a.csv"
    p.write_text("﻿name , email\nAna, ana@x.example \n,\n\nBo,bo@x.example\n", encoding="utf-8")
    headers, rows = read_rows(p)
    assert headers == ["name", "email"]
    assert rows == [{"name": "Ana", "email": "ana@x.example"}, {"name": "Bo", "email": "bo@x.example"}]


@pytest.mark.parametrize(
    "data",
    [
        "name;email\nJosé;jose@x.example\n".encode("cp1252"),  # Excel "CSV" in many countries
        "name\temail\nJosé\tjose@x.example\n".encode(),
        "name,email\nJosé,jose@x.example\n".encode("utf-8-sig"),
    ],
)
def test_read_csv_encodings_and_separators(tmp_path, data):
    p = tmp_path / "a.csv"
    p.write_bytes(data)
    assert read_rows(p) == (["name", "email"], [{"name": "José", "email": "jose@x.example"}])


def test_read_xlsx_named_sheet(tmp_path):
    wb = Workbook()
    wb.active.title = "Notes"
    ws = wb.create_sheet("Leads")
    ws.append(["name", "qty"])
    ws.append(["Ana", 3])
    ws.append([None, None])
    p = tmp_path / "a.xlsx"
    wb.save(p)
    headers, rows = read_rows(p, "Leads")
    assert headers == ["name", "qty"] and rows == [{"name": "Ana", "qty": 3}]
    with pytest.raises(JobError, match="Sheet 'Nope' not found"):
        read_rows(p, "Nope")


def test_write_all_formats(tmp_path):
    rows = [{"name": "Ana", "status": "done"}, {"name": "Bo", "status": "failed"}]
    cols = ["name", "status"]
    assert json.loads(write_rows(tmp_path / "r.json", cols, rows).read_text()) == rows
    assert write_rows(tmp_path / "r.csv", cols, rows).read_text().splitlines() == [
        "name,status",
        "Ana,done",
        "Bo,failed",
    ]
    ws = load_workbook(write_rows(tmp_path / "r.xlsx", cols, rows)).active
    assert [c.value for c in ws[1]] == cols
    assert ws["B2"].fill.fgColor.rgb.endswith("D1F2E3") and ws["B3"].fill.fgColor.rgb.endswith("F8D7DA")
    with pytest.raises(JobError):
        write_rows(tmp_path / "r.txt", cols, rows)
    # JSON has the same columns as the other formats
    extra = [{"name": "Ana", "status": "done", "internal": 1}]
    assert json.loads(write_rows(tmp_path / "e.json", cols, extra).read_text()) == [{"name": "Ana", "status": "done"}]
