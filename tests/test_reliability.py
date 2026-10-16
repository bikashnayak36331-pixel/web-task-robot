"""Regression tests for problems found in the 16 Oct 2026 end-to-end test (v1.1.0)."""

import json
import os
import signal
import subprocess
import sys
import time
from datetime import date, datetime

import pytest
from openpyxl import Workbook, load_workbook

from web_task_robot.cli import main
from web_task_robot.config import JobError, load_job
from web_task_robot.demo_site import DemoSite
from web_task_robot.runner import _key, run_extract, run_fill

from .conftest import EXAMPLES
from .lab import Lab


def job_file(tmp_path, text, rows=None, name="job.yaml"):
    if rows is not None:
        (tmp_path / "rows.csv").write_text(rows, encoding="utf-8")
    (tmp_path / name).write_text(text, encoding="utf-8")
    return load_job(tmp_path / name)


def slow_demo_site(delay):
    """The demo site, but every form post takes `delay` seconds, like a busy real server."""
    site = DemoSite(flaky=0)
    handler = site.httpd.RequestHandlerClass
    original = handler.do_POST

    def do_post(self):
        if self.path.startswith("/customers/new"):
            time.sleep(delay)
        return original(self)

    handler.do_POST = do_post
    return site


SLOW_JOB = """name: slow
type: fill
input: rows.csv
key: email
login:
  url: ${DEMO_URL}/login
  steps:
    - fill: { selector: "#email", value: "${PORTAL_USER}" }
    - fill: { selector: "#password", value: "${PORTAL_PASSWORD}" }
    - click: "button[type=submit]"
steps:
  - goto: ${DEMO_URL}/customers/new
  - fill: { selector: "#name", value: "{name}" }
  - fill: { selector: "#email", value: "{email}" }
  - select: { selector: "#plan", value: "{plan}" }
  - click: "#save"
success: ".alert-success"
error: ".alert-error"
options: { retries: 2, retry_delay_ms: 50, delay_ms: 0, timeout_ms: 1500 }
"""


def test_a_form_sent_to_a_slow_site_is_not_sent_again(tmp_path, monkeypatch, quiet):
    log, _ = quiet
    with slow_demo_site(2.5) as site:
        monkeypatch.setenv("DEMO_URL", site.url)
        monkeypatch.setenv("PORTAL_USER", site.state.user)
        monkeypatch.setenv("PORTAL_PASSWORD", site.state.password)
        job = job_file(tmp_path, SLOW_JOB, "name,email,plan\nNew Person,new@person.example,Pro\n")
        res = run_fill(job, tmp_path / "out" / "r", progress=log)
        time.sleep(1.5)  # let the slow post finish on the server
        posts = site.state.posts
    row = res.rows[0]
    assert row["status"] == "failed" and row["attempts"] == 1
    assert "not retried: the form may already have been sent" in row["message"]
    assert posts == 1, "v1.0.0 posted this form 3 times"


def test_a_server_error_after_sending_is_still_retried(tmp_path, monkeypatch, quiet):
    log, _ = quiet
    with DemoSite(flaky=0.5, seed=3) as site:
        monkeypatch.setenv("DEMO_URL", site.url)
        monkeypatch.setenv("PORTAL_USER", site.state.user)
        monkeypatch.setenv("PORTAL_PASSWORD", site.state.password)
        rows = "name,email,plan\n" + "".join(f"P{i},p{i}@x.example,Pro\n" for i in range(6))
        job = job_file(tmp_path, SLOW_JOB.replace("retries: 2", "retries: 8"), rows)
        res = run_fill(job, tmp_path / "out" / "r", progress=log)
    assert res.count("done") == 6 and any(r["attempts"] > 1 for r in res.rows)


SPA_JOB = """type: fill
input: rows.csv
key: email
steps:
  - goto: ${LAB}/spa-form
  - fill: { selector: "#name", value: "{name}" }
  - fill: { selector: "#email", value: "{email}" }
  - click: "#save"
  - extract: { selector: ".ref, .alert-error", as: answer }
success: ".alert-success"
error: ".alert-error"
options: { retries: 2, retry_delay_ms: 50, delay_ms: 0, timeout_ms: 3000 }
"""


def test_javascript_form_error_is_reported_once(tmp_path, monkeypatch, quiet):
    log, _ = quiet
    with Lab(delay=0.4) as lab:
        monkeypatch.setenv("LAB", lab.url)
        job = job_file(tmp_path, SPA_JOB, "name,email\nA,a@x.example\nB,not-an-email\n")
        res = run_fill(job, tmp_path / "out" / "r", progress=log)
        saved = [s["email"] for s in lab.saved]
    good, bad = res.rows
    assert good["status"] == "done" and good["answer"] == "R-1"
    assert bad["status"] == "failed" and bad["attempts"] == 1
    assert bad["message"] == "Site said: Email invalid or already used: not-an-email"
    assert saved == ["a@x.example"]


def test_javascript_paging_reads_each_page_once(tmp_path, monkeypatch, quiet):
    log, _ = quiet
    with Lab() as lab:
        monkeypatch.setenv("LAB", lab.url)
        job = job_file(
            tmp_path,
            "type: extract\nstart_url: ${LAB}/spa-list\nitem: li.it\nfields: {text: '.'}\n"
            "next_page: '#next:not([disabled])'\nmax_pages: 5\noutput: list.json\noptions: {delay_ms: 0}\n",
        )
        res = run_extract(job, tmp_path / "out" / "e", progress=log)
    texts = [r["text"] for r in json.loads(res.output.read_text())]
    assert texts == [f"Item {i}" for i in range(1, 26)], "v1.0.0 read 50 rows with only 20 different items"


def test_next_link_that_changes_nothing_stops_cleanly(tmp_path, monkeypatch, quiet):
    log, _ = quiet
    with Lab() as lab:
        monkeypatch.setenv("LAB", lab.url)
        job = job_file(
            tmp_path,
            "type: extract\nstart_url: ${LAB}/weird\nitem: .it\nfields: {t: .t}\nnext_page: .it\n"
            "max_pages: 5\noptions: {delay_ms: 0, timeout_ms: 1000}\n",
        )
        res = run_extract(job, tmp_path / "out" / "e", progress=log)
    assert len(res.rows) == 2 and "did not show new items" in res.note


def test_extract_with_no_items_says_so(tmp_path, monkeypatch, quiet):
    log, _ = quiet
    with Lab() as lab:
        monkeypatch.setenv("LAB", lab.url)
        job = job_file(
            tmp_path,
            "type: extract\nstart_url: ${LAB}/empty\nitem: .card\nfields: {t: .t}\noptions: {timeout_ms: 1000}\n",
        )
        with pytest.raises(JobError, match="no items matching '.card' appeared within 1s"):
            run_extract(job, tmp_path / "out" / "e", progress=log)
        assert main(["run", str(tmp_path / "job.yaml"), "--out", str(tmp_path / "out")]) == 2


def test_excel_dates_radios_uploads_and_dialogs(tmp_path, monkeypatch, quiet):
    log, _ = quiet
    wb = Workbook()
    ws = wb.active
    ws.title = "2024"
    ws.append(["id", "when", "size", "gift", "color", "qty", "doc"])
    ws.append([1001.0, datetime(2026, 10, 20), "M", "yes", "Green", 3.0, "a.txt"])
    ws.append([1002, date(2026, 11, 1), "S", "no", "r", 2, "a.txt"])
    wb.save(tmp_path / "orders.xlsx")
    (tmp_path / "a.txt").write_text("hello")
    text = """type: fill
input: orders.xlsx
sheet: 2024
key: id
steps:
  - goto: ${LAB}/widgets
  - click: "#ask"
  - extract: { selector: "#answer", as: answer }
  - fill: { selector: "#when", value: "{when}" }
  - check: { selector: "#size-{size}" }
  - check: { selector: "#size-S", when: "{gift}" }
  - check: { selector: "#gift", when: "{gift}" }
  - select: { selector: "#color", value: "{color}" }
  - fill: { selector: "#qty", value: "{qty}" }
  - upload: { selector: "#doc", file: "{doc}" }
  - click: "#go"
  - extract: { selector: "#echo", as: echo }
options: { retries: 0, delay_ms: 0, timeout_ms: 3000, dialogs: accept }
"""
    with Lab() as lab:
        monkeypatch.setenv("LAB", lab.url)
        res = run_fill(job_file(tmp_path, text), tmp_path / "out" / "w", progress=log)
        dismissed = run_fill(
            job_file(tmp_path, text.replace("dialogs: accept", "dialogs: dismiss"), name="j2.yaml"),
            tmp_path / "out2" / "w",
            progress=log,
        )
    first, second = res.rows
    assert [r["status"] for r in res.rows] == ["done", "done"], [r["message"] for r in res.rows]
    assert json.loads(first["echo"]) == {
        "color": ["g"],
        "doc": ["a.txt"],
        "gift": ["yes"],
        "qty": ["3"],
        "size": ["S"],  # "M", then S because gift is yes
        "when": ["2026-10-20"],  # v1.0.0 typed "2026-10-20 00:00:00", which a date field refuses
    }
    assert json.loads(second["echo"])["size"] == ["S"] and json.loads(second["echo"])["when"] == ["2026-11-01"]
    assert first["answer"] == "yes" and dismissed.rows[0]["answer"] == "no"
    assert "1001" in (tmp_path / "out" / "checkpoint.json").read_text()


def test_missing_upload_file_fails_the_row_without_retries(tmp_path, monkeypatch, quiet):
    log, _ = quiet
    with Lab() as lab:
        monkeypatch.setenv("LAB", lab.url)
        job = job_file(
            tmp_path,
            "type: fill\ninput: rows.csv\nsteps:\n  - goto: ${LAB}/widgets\n"
            "  - upload: {selector: '#doc', file: '{doc}'}\noptions: {delay_ms: 0}\n",
            "doc\nmissing.pdf\n",
        )
        row = run_fill(job, tmp_path / "out" / "u", progress=log).rows[0]
    assert row["status"] == "failed" and row["attempts"] == 1 and "File to upload not found" in row["message"]


def test_scraped_text_is_safe_in_excel(tmp_path, monkeypatch, quiet):
    log, _ = quiet
    with Lab() as lab:
        monkeypatch.setenv("LAB", lab.url)
        job = job_file(tmp_path, "type: extract\nstart_url: ${LAB}/weird\nitem: .it\nfields: {t: .t}\n")
        res = run_extract(job, tmp_path / "out" / "x", progress=log)
    cells = load_workbook(res.output).active["A"][1:]
    assert [(c.value, c.data_type) for c in cells] == [
        ('=HYPERLINK("http://evil.example","click")', "s"),  # text, not a live formula
        ("belland vt", "s"),  # control characters removed; v1.0.0 crashed here and lost every row
    ]


def test_blank_keys_and_screenshot_names(tmp_path, quiet):
    log, _ = quiet
    job = job_file(
        tmp_path,
        "type: fill\ninput: rows.csv\nkey: email\nsteps:\n  - goto: about:blank\n  - screenshot: '{name}'\n"
        "options: {delay_ms: 0}\n",
        "name,email\n../../escape,\nA/B,\nC,c@x.example\nD,c@x.example\n",
    )
    res = run_fill(job, tmp_path / "out" / "r", progress=log)
    assert [r["status"] for r in res.rows] == ["done", "done", "done", "skipped"]
    shots = sorted(p.name for p in (tmp_path / "out" / "r" / "screenshots").iterdir())
    assert shots == ["A_B.png", "C.png", "_.._escape.png"]
    assert not list(tmp_path.glob("*.png")) and not list((tmp_path / "out").glob("*.png"))


@pytest.mark.skipif(sys.platform == "win32", reason="POSIX signals")
def test_ctrl_c_saves_results_and_stops(tmp_path):
    with DemoSite(flaky=0) as site:
        env = dict(os.environ, DEMO_URL=site.url, PORTAL_USER=site.state.user, PORTAL_PASSWORD=site.state.password)
        for f in ("add-customers.yaml", "customers.csv"):
            (tmp_path / f).write_bytes((EXAMPLES / f).read_bytes())
        cmd = [sys.executable, "-m", "web_task_robot", "run", str(tmp_path / "add-customers.yaml")]
        p = subprocess.Popen(
            cmd + ["--out", str(tmp_path / "out"), "--slow-mo", "100"],
            env=env,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            start_new_session=True,  # like a terminal: Ctrl+C goes to every process in the group
        )
        deadline = time.time() + 60
        while time.time() < deadline and not list((tmp_path / "out").rglob("checkpoint.json")):
            time.sleep(0.2)
        os.killpg(p.pid, signal.SIGINT)
        out = p.communicate(timeout=60)[0]
        customers = len(site.state.customers)
    assert p.returncode == 130, out
    assert "Stopped with Ctrl+C after" in out
    results = list((tmp_path / "out").rglob("results.xlsx"))
    assert results and load_workbook(results[0]).active.max_row >= 2
    done = json.loads(next((tmp_path / "out").rglob("checkpoint.json")).read_text())
    assert customers == 1 + len(done), "every saved customer is in the checkpoint"


def test_demo_uses_its_own_login(tmp_path, capsys, monkeypatch):
    monkeypatch.chdir(tmp_path)
    monkeypatch.setenv("PORTAL_USER", "me@real-company.example")
    monkeypatch.setenv("PORTAL_PASSWORD", "real-secret")
    assert main(["demo", "--flaky", "0"]) == 0
    assert "7 done, 2 failed, 1 skipped" in capsys.readouterr().out
    assert os.environ["PORTAL_USER"] == "me@real-company.example", "the user's settings are put back"


def test_cli_rejects_bad_numbers(capsys):
    for args in (["run", "x.yaml", "--limit", "0"], ["run", "x.yaml", "--slow-mo", "-5"], ["demo", "--flaky", "2"]):
        with pytest.raises(SystemExit) as e:
            main(args)
        assert e.value.code == 2


def test_key_names_are_forgiving():
    assert _key("enter") == "Enter" and _key("ENTER") == "Enter" and _key("esc") == "Escape"
    assert _key("ctrl+a") == "Control+a" and _key("Shift+Tab") == "Shift+Tab" and _key("a") == "a"
