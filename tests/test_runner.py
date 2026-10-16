"""End-to-end tests: a real headless Chromium against the local demo site."""

import json
import shutil

import pytest
from openpyxl import load_workbook

from web_task_robot.cli import main
from web_task_robot.config import JobError, load_job
from web_task_robot.demo_site import DemoSite
from web_task_robot.runner import run_extract, run_fill, run_job

from .conftest import EXAMPLES


def fill_job(tmp_path, **opts):
    for f in ("add-customers.yaml", "customers.csv"):
        shutil.copy(EXAMPLES / f, tmp_path / f)
    job = load_job(tmp_path / "add-customers.yaml")
    job.options.delay_ms = 0
    for k, v in opts.items():
        setattr(job.options, k, v)
    return job


def by_email(res):
    out = {}
    for r in res.rows:
        out.setdefault(r["email"], r)  # first row wins; the input has one duplicate on purpose
    return out


def test_fill_end_to_end(site, tmp_path, quiet):
    log, lines = quiet
    res = run_fill(fill_job(tmp_path), tmp_path / "out" / "run1", progress=log)
    rows = by_email(res)
    assert res.count("done") == 7 and res.count("failed") == 2 and res.count("skipped") == 1
    assert rows["maya@brightside.example"]["reference"] == "CUS-1002"
    assert rows["tom@elmandoak.example"]["message"].startswith("Site said: A customer with email")
    assert rows["tom@elmandoak.example"]["attempts"] == 1, "site errors are not retried"
    assert rows["grace@kimaccounting"]["screenshot"].endswith("-failed.png")
    assert (res.out_dir / rows["grace@kimaccounting"]["screenshot"]).exists()
    saved = {c["email"]: c for c in site.state.customers}
    assert saved["maya@brightside.example"]["newsletter"] is True
    assert saved["lucas@meyerplumbing.example"]["newsletter"] is False
    assert saved["sofia@rossibakery.example"]["notes"] == "Two locations"
    ws = load_workbook(res.output).active
    assert ws.max_row == 11 and [c.value for c in ws[1]][-6:] == [
        "reference",
        "status",
        "message",
        "attempts",
        "seconds",
        "screenshot",
    ]
    assert "Rows" in res.report.read_text() and any("skipped (duplicate of row 1)" in line for line in lines)


def test_resume_skips_rows_done_earlier(site, tmp_path, quiet):
    log, _ = quiet
    job = fill_job(tmp_path)
    run_fill(job, tmp_path / "out" / "run1", progress=log)
    checkpoint = json.loads((tmp_path / "out" / "checkpoint.json").read_text())
    assert checkpoint["input"] == "customers.csv" and checkpoint["sites"] == [site.url]
    done = checkpoint["done"]
    assert len(done) == 7 and done["elena@petrova.example"] == {"reference": "CUS-1008"}
    again = run_fill(job, tmp_path / "out" / "run2", progress=log)
    assert again.count("skipped") == 8 and again.count("failed") == 2 and again.count("done") == 0
    assert by_email(again)["maya@brightside.example"]["reference"] == "CUS-1002"  # carried over
    assert len(site.state.customers) == 8, "nothing was submitted twice"


def test_retries_ride_out_a_flaky_server(tmp_path, monkeypatch, quiet):
    log, lines = quiet
    with DemoSite(flaky=0.4, seed=3) as s:
        monkeypatch.setenv("DEMO_URL", s.url)
        monkeypatch.setenv("PORTAL_USER", s.state.user)
        monkeypatch.setenv("PORTAL_PASSWORD", s.state.password)
        res = run_fill(fill_job(tmp_path, retries=6, retry_delay_ms=50), tmp_path / "out" / "r", progress=log)
    assert res.count("done") == 7
    assert any(r["attempts"] > 1 for r in res.rows if r["status"] == "done")
    assert any("HTTP 503" in line and "retrying" in line for line in lines)


def test_browser_validation_is_reported_without_retries(tmp_path, quiet):
    log, _ = quiet
    (tmp_path / "rows.csv").write_text("email\nnot-an-email\n")
    page = "data:text/html,<form action='/x'><input id='mail' type='email' required><button id='go'>Go</button></form>"
    (tmp_path / "job.yaml").write_text(
        f'type: fill\ninput: rows.csv\nsteps:\n  - goto: "{page}"\n'
        "  - fill: {selector: '#mail', value: '{email}'}\n  - click: '#go'\nsuccess: '.done'\n"
        "options: {timeout_ms: 3000, retries: 2}\n"
    )
    res = run_fill(load_job(tmp_path / "job.yaml"), tmp_path / "out" / "v", progress=log)
    row = res.rows[0]
    assert row["status"] == "failed" and row["attempts"] == 1
    assert row["message"].startswith("Form not accepted by the browser: #mail:")


def test_login_failure_stops_the_job(site, tmp_path, monkeypatch, quiet):
    log, _ = quiet
    monkeypatch.setenv("PORTAL_PASSWORD", "wrong")
    job = fill_job(tmp_path, timeout_ms=2000)
    with pytest.raises(JobError, match="Login failed"):
        run_fill(job, tmp_path / "out" / "x", progress=log)


def test_extract_all_pages(site, tmp_path, quiet):
    log, _ = quiet
    job = load_job(EXAMPLES / "collect-products.yaml")
    job.options.delay_ms = 0
    job.output = tmp_path / "products.json"
    res = run_extract(job, tmp_path / "out" / "e", progress=log)
    rows = json.loads(res.output.read_text())
    assert len(rows) == 30 and [r["page"] for r in rows].count(4) == 6
    first = rows[0]
    expected = site.state.products[0]
    assert first["sku"] == expected["sku"] and first["name"] == expected["name"]
    assert first["price"] == expected["price"] and isinstance(first["rating"], float)
    assert first["url"] == f"{site.url}/products/{expected['sku']}"
    job.max_pages = 2
    job.output = None
    assert len(run_extract(job, tmp_path / "out" / "e2", progress=log).rows) == 16


def test_cli_check_examples_and_run(site, tmp_path, capsys, monkeypatch):
    assert main(["examples", str(tmp_path / "jobs")]) == 0
    assert (tmp_path / "jobs" / "customers.csv").exists()
    assert main(["check", str(tmp_path / "jobs" / "add-customers.yaml")]) == 0
    assert "10 rows" in capsys.readouterr().out
    monkeypatch.delenv("PORTAL_PASSWORD")
    assert main(["check", str(tmp_path / "jobs" / "add-customers.yaml")]) == 2
    assert "PORTAL_PASSWORD" in capsys.readouterr().err
    monkeypatch.setenv("PORTAL_PASSWORD", site.state.password)
    code = main(["run", str(tmp_path / "jobs" / "add-customers.yaml"), "--out", str(tmp_path / "o"), "--limit", "3"])
    assert code == 0  # first three rows all succeed
    assert len(site.state.customers) == 4
    res = run_job(load_job(tmp_path / "jobs" / "collect-products.yaml"), tmp_path / "o")
    assert res.ok and res.output.name == "results.xlsx"


def test_demo_command(tmp_path, capsys, monkeypatch):
    monkeypatch.chdir(tmp_path)
    assert main(["demo", "--flaky", "0"]) == 0
    out = capsys.readouterr().out
    assert "Add customers to the portal: 7 done, 2 failed, 1 skipped" in out
    assert "Collect products: 30 done" in out and "now has 8 customers" in out
